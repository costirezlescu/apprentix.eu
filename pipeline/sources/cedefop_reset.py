"""Cedefop — Regional skills ecosystem index (RESET).

One /files/ download (like cedefop_skills.py): Cedefop's HTML pages refuse automated
requests, but the file link works. If the download fails, the most recent stored copy
in data/raw/cedefop/ is used. Licence CC BY 4.0, as declared for the dataset in
Cedefop's data.europa.eu catalogue (the workbook itself says "Copyright Cedefop").

The workbook has one row per EU NUTS 2 region (244), one per member state and one for
the EU-27, with the RESET score, its 4 pillars, 9 sub-pillars and 37 indicators, all
standardised from 1 to 100 (higher = more favourable).

Produces:
  indicators/cedefop-reset                  country and EU-27 scores (index, pillars, sub-pillars, indicators)
  regional-skills-ecosystems (records)      one record per NUTS 2 region, with its ranks
"""

from __future__ import annotations

import io
import re
import warnings

from ..common import (FetchError, PUBLISHED, countries, fetch, latest_raw, provenance, read_json, rel,
                      save_raw, write_indicator, write_json, write_records)

DATASET_ID = "regional-skills-ecosystems"
LANDING = "https://www.cedefop.europa.eu/en/datasets/regional-skills-ecosystem-index"
TOOL = "https://www.cedefop.europa.eu/en/tools/reset"
FILE = {"file": "dataset-regional-skills-ecosystem-index-reset.xlsx", "raw_name": "reset.xlsx"}
URL = "https://www.cedefop.europa.eu/files/" + FILE["file"]
LICENCE_NOTE = ("Licence CC BY 4.0 as declared for the dataset in Cedefop's data.europa.eu catalogue "
                "(the workbook states: Copyright Cedefop, 2026).")

SOURCE = {
    "id": "cedefop-reset",
    "name": "Cedefop — Regional skills ecosystem index (RESET)",
    "publisher": "Cedefop",
    "homepage": LANDING,
    "description": "Composite index comparing the skills ecosystems of the EU's NUTS 2 regions: skills supply, demand, "
                   "match and environment, with VET-specific indicators (tasks, activity, employment and unemployment "
                   "of people with initial VET). Published as country scores and a regional dataset.",
    "access": "file",
    "browser_cors": False,
    "licence": "CC-BY-4.0",
    "cadence": "new in 2026; update cycle not yet known",
    "secret": None,
    "outputs": ["indicators/cedefop-reset", DATASET_ID],
}

warnings.filterwarnings("ignore", message="Unknown extension is not supported", module="openpyxl")

KIND_PREFIX = {"Index": "", "Pillar": "Pillar: ", "Subpillar": "Sub-pillar: ", "Indicator": "Indicator: "}
# VET-specific indicators shown in each region's record ("(IVET)" / "(secondary)" in the workbook).
VET_INDICATORS = ["cog_task_ivet", "dig_task_ivet", "act_you_ivet", "act_adu_ivet", "emp_rec_ivet",
                  "unemp_ivet", "unemplong_sec", "cvet_sec"]
BANDS = ["Top fifth of EU regions", "Second fifth", "Middle fifth", "Fourth fifth", "Bottom fifth of EU regions"]


def download() -> tuple[bytes, object, object]:
    try:
        body = fetch(URL, timeout=180)
        if body[:2] == b"PK":
            raw, digest = save_raw("cedefop", FILE["raw_name"], body)
            return body, raw, digest
    except FetchError:
        pass
    raw = latest_raw("cedefop", "*-reset.xlsx")
    if not raw:
        raise RuntimeError(f"RESET download failed and no stored copy exists ({URL})")
    import hashlib
    body = raw.read_bytes()
    return body, raw, hashlib.sha256(body).hexdigest()


def _s(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def parse(body: bytes) -> dict:
    """{'readme': {...}, 'columns': [(key, label, kind)], 'rows': [{country, code, name, name_latin, values{}}]}"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    readme = {_s(r[0]).rstrip(":").lower(): r[1] for r in wb["Readme"].iter_rows(values_only=True) if r and r[0]}
    rows = list(wb["data"].iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and _s(r[0]) == "Country" and _s(r[1]) == "Region")
    kinds = next(r for r in rows[:hdr_i] if r and _s(r[3]) == "Value type")
    names = next(r for r in rows[:hdr_i] if r and _s(r[3]) == "Display name")
    keys = rows[hdr_i]
    if _s(keys[4]) != "RESET":
        raise RuntimeError(f"RESET workbook layout changed: expected 'RESET' in column E, found {keys[4]!r}")
    columns = []
    for i in range(4, len(keys)):
        k = _s(keys[i])
        if not k:
            continue
        kind = _s(kinds[i]) or "Index"
        columns.append((i, k, _s(names[i]) or k, kind))
    out = []
    for r in rows[hdr_i + 1:]:
        if not r or not _s(r[1]):
            continue
        vals = {k: float(r[i]) for i, k, _, _ in columns if isinstance(r[i], (int, float))}
        out.append({"country": _s(r[0]), "code": _s(r[1]), "name": _s(r[2]), "name_latin": _s(r[3]), "values": vals})
    if len(out) < 200:
        raise RuntimeError(f"RESET workbook has only {len(out)} rows; refusing to publish a partial dataset")
    return {"readme": readme, "columns": [(k, label, kind) for _, k, label, kind in columns], "rows": out}


def _version(readme: dict) -> tuple[str, str]:
    version = _s(readme.get("version")) or "2026"
    year = _s(readme.get("reference year(s)") or readme.get("reference year")) or "2024"
    return version, re.sub(r"\.0$", "", year)


def indicator(p: dict, raw, digest) -> str:
    version, year = _version(p["readme"])
    labels = {k: KIND_PREFIX.get(kind, "") + (label if kind != "Index" else "RESET score (overall)")
              for k, label, kind in p["columns"]}
    series = []
    for r in p["rows"]:
        geo = "EU27" if r["code"] == "EU27" else r["code"] if len(r["code"]) == 2 else None
        if not geo:
            continue
        for k, v in r["values"].items():
            series.append({"geo": geo, "time": year, "value": round(v, 2), "dims": {"component": k}})
    ind = {
        "id": "cedefop-reset",
        "title": "Regional skills ecosystem index (RESET), country scores",
        "description": (
            "Cedefop's RESET compares how well skills are developed, used and supported, combining indicators in "
            "four pillars: skills supply, skills demand, skills match and the skills environment (quality of "
            "education and business environment, education–business cooperation). Several indicators look "
            "specifically at people with initial VET (IVET): the cognitive and digital content of their jobs, "
            "their activity, employment and unemployment. Scores run from 1 to 100; higher is more favourable. "
            "These are the country and EU-27 rows of the dataset; the 244 regional scores are in the "
            "'Regional skills ecosystems' dataset."),
        "unit": "score (1–100)",
        "topic": "Skills intelligence",
        "source_label": f"Cedefop, Regional skills ecosystem index (RESET), {version} (data sheet)",
        "comparability": (
            f"Reference year {year}; one edition so far, so there is no time series. Scores are standardised "
            "(1–100) relative to the other EU regions, not raw values, and are designed to compare places with each "
            "other. RESET uses a different framework from Cedefop's European Skills Index (ESI); the two are not "
            "comparable. Covers EU member states only."),
        "dims": [{"key": "component", "label": "Index, pillar, sub-pillar or indicator",
                  "values": labels, "default": "RESET"}],
        "provenance": provenance(
            publisher="Cedefop", dataset_code=f"RESET {version}", source_url=LANDING, api_url=URL,
            licence="CC-BY-4.0",
            citation=f"Cedefop (2026). Regional skills ecosystem index (RESET), {version} [Data set]. {LANDING}. "
                     + LICENCE_NOTE,
            source_updated=version, raw=raw, raw_sha256=digest),
        "series": series,
    }
    return rel(write_indicator(ind))


def _fmt(v: float | None) -> str | None:
    return None if v is None else f"{v:.1f}"


def records(p: dict, raw, digest) -> list[str]:
    version, year = _version(p["readme"])
    ref = countries()["by_code"]
    labels = {k: label for k, label, _ in p["columns"]}
    kinds = {k: kind for k, _, kind in p["columns"]}
    regions = [r for r in p["rows"] if len(r["code"]) == 4 and r["code"] != "EU27" and "RESET" in r["values"]]
    ranked = sorted(regions, key=lambda r: -r["values"]["RESET"])
    eu_rank = {r["code"]: i + 1 for i, r in enumerate(ranked)}
    n = len(regions)
    by_country: dict[str, list] = {}
    for r in ranked:
        by_country.setdefault(r["code"][:2], []).append(r)
    pillars = [k for k in labels if kinds[k] == "Pillar"]
    subs = [k for k in labels if kinds[k] == "Subpillar"]
    vet = [k for k in VET_INDICATORS if k in labels]

    recs = []
    for r in sorted(regions, key=lambda r: (r["country"], -r["values"]["RESET"])):
        cc = r["code"][:2]
        c = ref.get(cc) or {}
        same = by_country[cc]
        rank = eu_rank[r["code"]]
        rec = {
            "id": f"reset-{r['code'].lower()}",
            "country": c.get("name") or r["country"],
            "country_code": cc,
            "flag": c.get("flag", ""),
            "region": r["name_latin"] or r["name"],
            "region_original": r["name"] if r["name"] and r["name"] != r["name_latin"] else None,
            "nuts2": r["code"],
            "reset": _fmt(r["values"]["RESET"]),
            "eu_rank": f"{rank} of {n}",
            "country_rank": f"{same.index(r) + 1} of {len(same)}" if len(same) > 1 else "Only region",
            "band": BANDS[min(4, (rank - 1) * 5 // n)],
            # Short labelled versions for the explorer's cards.
            "chip_score": f"Score {r['values']['RESET']:.1f} / 100",
            "chip_rank": f"#{rank} of {n} EU regions",
            "chip_country": (f"#{same.index(r) + 1} of {len(same)} in {c.get('name') or r['country']}"
                             if len(same) > 1 else f"Only region of {c.get('name') or r['country']}"),
        }
        for k in pillars + subs + vet:
            rec[f"s_{k}"] = _fmt(r["values"].get(k))
        recs.append({k: v for k, v in rec.items() if v is not None})

    def field(k, label, type_="text", **kw):
        return {"key": k, "label": label, "type": type_, **kw}
    fields = [
        field("country", "Country", "category", facet=True, collapse=12),
        field("band", "Position among EU regions", "category", facet=True),
        field("region", "Region", "title"),
        field("region_original", "Name in the national language", "subtitle"),
        field("nuts2", "NUTS 2 code"),
        field("reset", "RESET score (1–100)"),
        field("eu_rank", "Rank among EU regions"),
        field("country_rank", "Rank within the country"),
        *[field(f"s_{k}", labels[k].strip()) for k in pillars + subs + vet],
        field("chip_score", "Score (card)"),
        field("chip_rank", "EU rank (card)"),
        field("chip_country", "Country rank (card)"),
        field("country_code", "Country code", "hidden"),
        field("flag", "Flag", "hidden"),
        field("id", "Identifier", "hidden"),
    ]
    meta = {
        "id": DATASET_ID,
        "title": "Regional skills ecosystems (RESET)",
        "tagline": "How Europe's 244 regions compare on skills supply, demand, match and environment.",
        "description": (
            "Cedefop's Regional skills ecosystem index (RESET) scores each EU NUTS 2 region from 1 to 100 on how "
            "skills are developed, used and supported, in four pillars and nine sub-pillars. Each record is one "
            "region with its overall score, its rank among all EU regions and within its country, and its scores "
            "on the VET-specific indicators (people with initial VET: job tasks, activity, employment, "
            f"unemployment). Reference year {year}."),
        "recordLabel": {"one": "region", "many": "regions"},
        "source": {
            "name": "Cedefop — Regional skills ecosystem index (RESET)",
            "url": LANDING,
            "licence": "CC BY 4.0. " + LICENCE_NOTE + " Cite as: Cedefop (2026), Regional skills ecosystem index "
                       "(RESET); compiled by Apprentix.",
            "licence_id": "CC-BY-4.0",
            "caveat": ("Scores are standardised from 1 to 100 relative to the other EU regions (higher = more "
                       "favourable); they are not raw statistics. Ranks and the 'position among EU regions' bands are "
                       f"calculated by Apprentix from the overall score. Version {version}, reference year {year}. "
                       f"Cedefop's interactive tool: {TOOL}"),
            "retrieved": None,
            "file_url": URL,
            "raw_file": rel(raw),
            "raw_sha256": digest,
        },
        "display": {"title": "region", "subtitle": "region_original", "group": "country", "badge": "flag",
                    "link": None, "facts": ["chip_score", "chip_rank", "chip_country"]},
        "sections": [
            {"title": "At a glance", "fields": ["country", "nuts2", "reset", "eu_rank", "country_rank", "band"]},
            {"title": "Pillars", "fields": [f"s_{k}" for k in pillars]},
            {"title": "Sub-pillars", "fields": [f"s_{k}" for k in subs]},
            {"title": "People with VET (IVET or secondary education)", "fields": [f"s_{k}" for k in vet]},
        ],
        "fields": fields,
    }
    meta["display"] = {k: v for k, v in meta["display"].items() if v is not None}

    from ..common import today
    folder = PUBLISHED / DATASET_ID
    old_rec, old_meta = folder / "records.json", folder / "meta.json"
    meta["source"]["retrieved"] = today()
    if old_rec.exists() and old_meta.exists() and read_json(old_rec) == recs:
        prev = read_json(old_meta).get("source", {})
        for k in ("retrieved", "raw_file", "raw_sha256"):
            if k in prev:
                meta["source"][k] = prev[k]
    rec_path = write_records(DATASET_ID, recs)
    write_json(old_meta, meta)
    return [rel(rec_path), rel(old_meta)]


def run() -> list[str]:
    body, raw, digest = download()
    p = parse(body)
    return [indicator(p, raw, digest), *records(p, raw, digest)]
