"""Cedefop downloadable datasets (the /files/ downloads in Cedefop's Datasets section).

Cedefop's web pages refuse automated requests, but files under /files/ download
normally. File names change with each release, so the current URLs are listed in
FILES below; the 'cedefop-watch' connector flags when a new release appears.
If a download is refused, the most recent copy in data/raw/cedefop/ is used.

Produces:
  indicators/cedefop-kivet-<number>   58 Key indicators on VET (CC BY 4.0)
  vet-policy-timeline/records.json    Timeline of VET policies in Europe (CC BY 4.0)
"""

from __future__ import annotations

import io
import re

from ..common import (FetchError, PUBLISHED, fetch, geo_code, html_to_text, latest_raw,
                      num, provenance, read_json, rel, save_raw, write_indicator,
                      write_json, write_records)

SOURCE = {
    "id": "cedefop",
    "name": "Cedefop — Datasets (Key indicators on VET; Timeline of VET policies)",
    "publisher": "Cedefop",
    "homepage": "https://www.cedefop.europa.eu/en/datasets",
    "description": "Cedefop's downloadable datasets: 58 Key indicators on VET for 36 countries (2010 onwards) and more than 1,000 national VET policy developments since 2015.",
    "access": "file",
    "browser_cors": False,
    "licence": "CC-BY-4.0",
    "cadence": "annual (KIVET June–July; Timeline yearly)",
    "secret": None,
}

FILES = {
    "kivet": {
        "url": "https://www.cedefop.europa.eu/files/integratedmasterfileIndicators_2026_%2006%20July.xlsx",
        "raw_name": "kivet-masterfile.xlsx",
        "raw_glob": "*-kivet-masterfile*.xlsx",
        "landing": "https://www.cedefop.europa.eu/en/datasets/dataset-key-indicators-vet",
        "doi": "10.2906/534056862890980",
        "citation": "Cedefop (2026). Key indicators on VET [Data set]. https://doi.org/10.2906/534056862890980",
        "version": "2026 (6 July 2026)",
    },
    "timeline": {
        "url": "https://www.cedefop.europa.eu/files/cedefop_timeline_of_vet_policies_in_europe_dataset_{year}.xlsx",
        "year": 2026,
        "raw_name": "timeline-of-vet-policies.xlsx",
        "raw_glob": "*-timeline-of-vet-policies*.xlsx",
        "landing": "https://www.cedefop.europa.eu/en/datasets/timeline-EU-vet-policies",
        "citation": "Cedefop, & ReferNet. (2026). Timeline of VET policies in Europe (2025 update): Research dataset [Data set]. https://www.cedefop.europa.eu/en/datasets/timeline-EU-vet-policies",
    },
}


def download(key: str) -> tuple[bytes, str, object]:
    """Fetch a Cedefop file, falling back to the latest stored original."""
    f = FILES[key]
    url = f["url"]
    if key == "timeline":
        # Try next year's file first: the name only changes by year.
        for year in (f["year"] + 1, f["year"]):
            candidate = url.format(year=year)
            try:
                body = fetch(candidate, timeout=180)
                if body[:2] == b"PK":  # xlsx is a zip
                    raw, digest = save_raw("cedefop", f["raw_name"], body)
                    return body, candidate, (raw, digest)
            except FetchError:
                continue
        url = url.format(year=f["year"])
    else:
        try:
            body = fetch(url, timeout=180)
            if body[:2] == b"PK":
                raw, digest = save_raw("cedefop", f["raw_name"], body)
                return body, url, (raw, digest)
        except FetchError:
            pass
    path = latest_raw("cedefop", f["raw_glob"])
    if not path:
        raise FetchError(f"Could not download {url} and no stored copy in data/raw/cedefop/")
    print(f"  using stored copy {rel(path)}")
    body = path.read_bytes()
    from ..common import sha256
    return body, url, (path, sha256(body))


# ------------------------------------------------------------- KIVET -------

def _cell(v):
    return None if v is None else str(v).strip()


def parse_kivet(body: bytes) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)

    # Indicator list on the ReadMe sheet.
    catalogue = {}
    for row in wb["ReadMe"].iter_rows(values_only=True):
        if row and isinstance(row[0], str) and re.fullmatch(r"Ind\w+", row[0].strip()):
            (link, no, label, desc, base, last, position, grp, typ, dgrp, dtyp) = (list(row) + [None] * 11)[:11]
            catalogue[link.strip()] = {
                "number": _cell(no), "label": _cell(label), "description": _cell(desc),
                "position": _cell(position), "group": _cell(grp), "type": _cell(typ),
                "dashboard_group": _cell(dgrp), "dashboard_type": _cell(dtyp),
            }

    out = []
    for sheet, info in catalogue.items():
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        source = None
        start = None
        for i, r in enumerate(rows):
            first = next((c for c in r if c is not None and str(c).strip()), None)
            if first == "Source" and source is None and i + 1 < len(rows):
                source = next((str(c).strip() for c in rows[i + 1] if c), None)
            if isinstance(first, str) and first.startswith("Indicator values and flags"):
                start = i
                break
        if start is None:
            continue
        # Year header row, then Value/Flag row, then one row per geography.
        hdr_i = next(i for i in range(start + 1, len(rows)) if any(
            isinstance(c, (str, int)) and re.fullmatch(r"(19|20)\d\d", str(c).strip() or "") for c in rows[i]))
        hdr = rows[hdr_i]
        year_cols = [(j, str(c).strip()) for j, c in enumerate(hdr)
                     if c is not None and re.fullmatch(r"(19|20)\d\d", str(c).strip())]
        series = []
        for r in rows[hdr_i + 2:]:
            code = _cell(r[1]) if len(r) > 1 else None
            if not code:
                break
            geo = geo_code(code)
            if not geo:
                continue
            for j, year in year_cols:
                v = num(r[j]) if j < len(r) else None
                flag = _cell(r[j + 1]) if j + 1 < len(r) else None
                if v is not None:
                    series.append({"geo": geo, "time": year, "value": v,
                                   **({"flag": flag} if flag else {})})
        label = info["label"] or sheet
        unit = "%" if "%" in label or "(%)" in label else ("1000s" if "1000s" in label else "")
        if "PPS" in label:
            unit = "1000 PPS"
        if "EUR" in label:
            unit = "EUR"
        if "per 1000 hours" in label:
            unit = "hours per 1000 hours worked"
        if "foreign languages" in label:
            unit = "languages"
        out.append({"sheet": sheet, "info": info, "source": source, "series": series, "unit": unit})
    return out


KIVET_FLAGS = {
    "b": "break in time series", "d": "definition differs", "e": "estimated",
    "p": "provisional", "u": "low reliability", "z": "not applicable", "c": "confidential",
}


def run_kivet() -> list[str]:
    body, url, (raw, digest) = download("kivet")
    f = FILES["kivet"]
    written = []
    for k in parse_kivet(body):
        info = k["info"]
        num_id = info["number"].lower()
        ind = {
            "id": f"cedefop-kivet-{num_id}",
            "title": info["label"],
            "description": info["description"],
            "unit": k["unit"],
            "topic": info["group"] or "Key indicators on VET",
            "source_label": f"Cedefop Key indicators on VET, indicator {info['number']}",
            "comparability": "EU averages are Cedefop estimates (weighted averages of available country data) where flagged.",
            "flags": KIVET_FLAGS,
            "kivet": {k2: v for k2, v in info.items() if v and k2 not in ("label", "description")},
            "provenance": provenance(
                publisher="Cedefop", dataset_code=f"KIVET {info['number']}",
                source_url=f["landing"], api_url=url, licence="CC-BY-4.0",
                citation=f["citation"] + (f" Underlying source: {k['source']}" if k["source"] else ""),
                source_updated=f["version"], raw=raw, raw_sha256=digest),
            "series": k["series"],
        }
        written.append(rel(write_indicator({k2: v for k2, v in ind.items() if v not in (None, "", {})})))
    return written


# ------------------------------------------------------------- Timeline ----

STAGE_ORDER = ["Design phase", "Pilot phase", "Legislative process", "Approved/agreed",
               "Implementation phase", "Completed", "Discontinued"]

COMMUNITY = {
    "BEFR": ("BE", "Belgium (French-speaking community)"),
    "BEFL": ("BE", "Belgium (Flemish community)"),
    "BENL": ("BE", "Belgium (Flemish community)"),
    "BEDE": ("BE", "Belgium (German-speaking community)"),
}


# Cedefop theme 3.9 is "Reinforcing work-based learning, including apprenticeships".
APPRENTICESHIP_TEXT = re.compile(
    r"apprentic|work-based learning|dual (vet|training|system|education|study|studies)|\bWBL\b", re.I)


def split(v) -> list[str]:
    if v is None:
        return []
    return [p.strip() for p in str(v).split("^") if p and p.strip()]


def parse_timeline(body: bytes) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    rows = list(wb["TIMELINE_DATA"].iter_rows(values_only=True))
    hdr = [(_cell(h) or "") for h in rows[0]]
    ix = {h: i for i, h in enumerate(hdr) if h}

    def col(r, name):
        i = ix.get(name)
        return r[i] if i is not None and i < len(r) else None

    stage_cols = sorted((h for h in ix if re.match(r"^20\d\d - Prog(r)?ess Stage$", h)), key=lambda h: h[:4])
    records = []
    for r in rows[1:]:
        pid = _cell(col(r, "ID"))
        if not pid:
            continue
        raw_country = _cell(col(r, "Country")) or ""
        cc, _, cname = raw_country.partition(" - ")
        if cc in COMMUNITY:
            code, cname = COMMUNITY[cc]
        else:
            code = geo_code(cc) or cc
        stages = []
        for h in stage_cols:
            s = _cell(col(r, h))
            if s:
                desc_col = next((d for d in ix if d.startswith(h[:4]) and "Description" in d), None)
                stages.append({"year": h[:4], "stage": s,
                               "description": html_to_text(col(r, desc_col)) if desc_col else None})
        latest = stages[-1] if stages else None
        themes = split(col(r, "Thematic categories"))
        top_themes = [t for t in themes if re.match(r"^\d+\. ", t)]
        sub_themes = [t for t in themes if re.match(r"^\d+\.\d+\. ", t)]
        targets = split(col(r, "Target groups"))
        text_blob = " ".join(filter(None, [
            _cell(col(r, "Title")),
            html_to_text(col(r, "Objectives of the policy development")),
            html_to_text(col(r, "Description of the policy development"))]))
        apprentices = (any(t.startswith("3.9.") for t in sub_themes)
                       or bool(APPRENTICESHIP_TEXT.search(text_blob)))
        urls = split(col(r, "Source(s) URL"))
        rec = {
            "id": f"tl-{pid}",
            "policy_id": pid,
            "title": html_to_text(col(r, "Title")),
            "country": cname or raw_country,
            "country_code": code,
            "type": _cell(col(r, "Type")),
            "subsystem": split(col(r, "Subsystem")),
            "concerns_apprenticeship": "Yes" if apprentices else "No",
            "target_groups": targets,
            "other_target_groups": html_to_text(col(r, "Other")),
            "themes": top_themes,
            "sub_themes": sub_themes,
            "eu_priorities": split(col(r, "EU priorities")),
            "reporting_period": [p.replace(" ", "") for p in split(col(r, "Reporting Period"))],
            "bodies": split(col(r, "Bodies")),
            "latest_stage": latest["stage"] if latest else None,
            "latest_year": latest["year"] if latest else None,
            "first_year": stages[0]["year"] if stages else None,
            "progress": " → ".join(f"{s['year']}: {s['stage']}" for s in stages) or None,
            "latest_stage_description": latest["description"] if latest else None,
            "background": html_to_text(col(r, "The WHY/Background of the policy development")),
            "objectives": html_to_text(col(r, "Objectives of the policy development")),
            "description": html_to_text(col(r, "Description of the policy development")),
            "related_ids": [f"tl-{x}" for x in split(col(r, "Related Policy Developments IDs"))],
            "source_titles": split(col(r, "Source(s) title")),
            "source_url": urls[0] if urls else None,
            "more_sources": urls[1:],
            "status": "machine-imported",
        }
        records.append({k: v for k, v in rec.items() if v not in (None, "", [])})
    records.sort(key=lambda x: (x["country"], -(int(x.get("latest_year") or 0)), x["title"] or ""))
    return records


def run_timeline() -> list[str]:
    body, url, (raw, digest) = download("timeline")
    f = FILES["timeline"]
    records = parse_timeline(body)
    for rec in records:
        c = rec.get("country_code")
        flag_code = {"EL": "GR", "UK": "GB"}.get(c, c)
        if flag_code and len(flag_code) == 2:
            rec["flag"] = "".join(chr(0x1F1E6 + ord(ch) - 65) for ch in flag_code)
    written = [rel(write_records("vet-policy-timeline", records))]
    meta_path = PUBLISHED / "vet-policy-timeline" / "meta.json"
    meta = read_json(meta_path)
    # The raw file name starts with the date it was first seen, so this only moves when data changes.
    meta["source"]["retrieved"] = raw.name[:10]
    meta["source"]["file_url"] = url
    meta["source"]["raw_file"] = rel(raw)
    meta["source"]["raw_sha256"] = digest
    write_json(meta_path, meta)
    written.append(rel(meta_path))
    return written


def run() -> list[str]:
    return run_kivet() + run_timeline()
