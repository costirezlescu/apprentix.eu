"""France — DARES apprenticeship contract series (new contracts and stock), via POEM.

POEM ("Politiques de l'emploi", https://poem.travail-emploi.gouv.fr) is DARES's
open-data portal for employment-policy schemes. Each series is downloadable as
XLSX/CSV/JSON/SDMX at /exports/donnees/<code>/<code>.<ext>; the list of series
is on https://poem.travail-emploi.gouv.fr/open-data/alternance-1 (used here for
discovery, so a renamed or withdrawn code fails loudly instead of silently).

The DARES site itself (dares.travail-emploi.gouv.fr, "Le contrat d'apprentissage")
sits behind an interactive bot check and cannot be fetched automatically; the
POEM series are the same DARES "Système d'information sur l'apprentissage" data.

InserJeunes (DEPP/DARES employment rates of former apprentices) is not included:
data.gouv.fr / data.education.gouv.fr publish it only as per-CFA rates without
leaver counts or a national row, so a correct national rate cannot be derived.
"""

from __future__ import annotations

import io
import re

from ..common import fetch, provenance, rel, save_raw, write_indicator

BASE = "https://poem.travail-emploi.gouv.fr"
CATALOGUE = BASE + "/open-data/alternance-1"
EXPORT = BASE + "/exports/donnees/{code}/{code}.xlsx"
PAGE = BASE + "/donnee/{slug}"

LICENCE = (
    "unverified — POEM legal notice (mentions légales) states the site's contents are "
    "\"des information publiques librement et gratuitement réutilisables dans les conditions "
    "fixées par la loi n°78-753 du 17 juillet 1978\", but also asks reusers to contact "
    "dares.dsypef at travail.gouv.fr; no named open licence is attached to the series"
)
PUBLISHER = "DARES (Ministère du Travail), POEM"

SOURCE = {
    "id": "fr-dares",
    "name": "DARES — apprenticeship contracts (POEM)",
    "publisher": PUBLISHER,
    "homepage": "https://poem.travail-emploi.gouv.fr/synthese/contrats-d-apprentissage",
    "description": "Monthly new apprenticeship contracts and end-of-month stock of apprentices in France since 2013, by employer sector (private/public) and level (secondary/higher education).",
    "access": "file",
    "browser_cors": None,
    "licence": LICENCE,
    "cadence": "monthly (about two months' lag)",
    "secret": None,
    "outputs": ["indicators/fr-dares-new-contracts", "indicators/fr-dares-apprentices-stock"],
}

MONTHS = {"janv": 1, "fév": 2, "fev": 2, "mars": 3, "avril": 4, "avr": 4, "mai": 5, "juin": 6,
          "juil": 7, "août": 8, "aout": 8, "sept": 9, "oct": 10, "nov": 11, "déc": 12, "dec": 12}

# (dimension code, label, POEM series suffix, POEM page slug)
BREAKDOWNS = [
    ("TOTAL", "All contracts (private + public sector)", "tot", "contrat-dapprentissage-secteurs-prive-et-public"),
    ("PRIVATE", "Private-sector employers", "prive", "contrat-dapprentissage-prive"),
    ("PUBLIC", "Public-sector employers", "public", "contrat-dapprentissage-public"),
    ("SECONDARY", "Secondary-level diplomas (CAP, Bac pro, BP…)", "second", "contrat-dapprentissage-enseignement-secondaire"),
    ("HIGHER", "Higher-education diplomas (BTS, licence, master, engineering…)", "sup", "contrat-dapprentissage-enseignement-superieur"),
]

COMMON_NOTE = (
    "France, national source (DARES apprenticeship information system: Ari@ne and consular "
    "chambers 2013–2019, Deca from 2020, with DARES estimates for late reporting — using DSN for "
    "the most recent year). Covers ALL apprenticeship contracts, including the large and growing "
    "share preparing higher-education diplomas (Bac+2 to master/engineer, i.e. EQF 5–7), which "
    "most other countries' apprenticeship counts do not include. Geography: all French "
    "départements including the overseas DROM, excluding COM; employer establishment address. "
    "Raw data, not seasonally adjusted (strong September peak). The 'Scope' breakdown gives two "
    "alternative partitions of the total (by employer sector, and by level), so do not add "
    "sector and level values together. Do not sum with other countries."
)

SPECS = [
    {
        "id": "fr-dares-new-contracts",
        "kind": "ent",
        "title": "New apprenticeship contracts in France (monthly)",
        "description": "Number of new apprenticeship contracts starting each month in France (including renewals), by employer sector and level of the diploma prepared.",
        "topic": "Participation",
        "unit": "contracts",
        "provisional_months": 36,
        "note": "Flow: new contracts by contract start month, including renewals (reconduction). "
                "DARES treats the last 36 months as provisional (flag 'p'). ",
    },
    {
        "id": "fr-dares-apprentices-stock",
        "kind": "stk",
        "title": "Apprentices under contract in France (end of month)",
        "description": "Number of apprentices with a current apprenticeship contract on the last day of each month in France, by employer sector and level of the diploma prepared.",
        "topic": "Participation",
        "unit": "apprentices",
        "provisional_months": 48,
        "note": "Stock: apprentices under contract on the last day of the month. DARES generally "
                "treats the last 48 months as provisional (flag 'p'). ",
    },
]


def _month(label: str) -> str | None:
    m = re.match(r"^(\d{4})-([^.\s]+)\.?$", str(label or "").strip())
    if not m:
        return None
    mo = MONTHS.get(m.group(2).lower())
    return f"{m.group(1)}-{mo:02d}" if mo else None


def _parse(body: bytes) -> tuple[dict, str | None, str | None]:
    """Return ({YYYY-MM: value} for France entière, last-update date, series title)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    info = list(wb.worksheets[0].iter_rows(values_only=True))
    title = info[0][0] if info else None
    updated = None
    for r in info:
        if r and r[0] and "mise à jour" in str(r[0]) and len(r) > 1 and r[1]:
            d = re.match(r"(\d{2})/(\d{2})/(\d{4})", str(r[1]))
            updated = f"{d.group(3)}-{d.group(2)}-{d.group(1)}" if d else str(r[1])
    rows = list(wb["National"].iter_rows(values_only=True))
    header = next(r for r in rows if r and r[0] == "Code_National")
    fre = next(r for r in rows if r and r[0] == "FRE")
    out = {}
    for h, v in zip(header[2:], fre[2:]):
        t = _month(h)
        if t is None:
            if h is not None:
                raise ValueError(f"POEM: unrecognised month header {h!r}")
            continue
        if isinstance(v, (int, float)):
            out[t] = v
    return out, updated, title


def _shift(month: str, n: int) -> str:
    y, m = map(int, month.split("-"))
    k = y * 12 + (m - 1) - n
    return f"{k // 12}-{k % 12 + 1:02d}"


def run() -> list[str]:
    catalogue = fetch(CATALOGUE).decode("utf-8", "replace")
    available = set(re.findall(r"/exports/donnees/([a-z0-9_]+)/", catalogue))
    written = []
    for spec in SPECS:
        series, updated, raws = [], [], []
        for code_dim, _, suffix, _ in BREAKDOWNS:
            code = f"new_cap_{suffix}_{spec['kind']}"
            if code not in available:
                raise ValueError(f"POEM series {code} not listed on {CATALOGUE}")
            body = fetch(EXPORT.format(code=code), timeout=120)
            raw, digest = save_raw("fr-dares", f"{code}.xlsx", body)
            raws.append((code, raw, digest))
            values, upd, _ = _parse(body)
            if upd:
                updated.append(upd)
            last = max(values)
            cutoff = _shift(last, spec["provisional_months"] - 1)
            for t, v in values.items():
                series.append({"geo": "FR", "time": t, "value": v, "dims": {"scope": code_dim},
                               **({"flag": "p"} if t >= cutoff else {})})
        total_code, total_raw, total_digest = raws[0]
        page_slug = BREAKDOWNS[0][3] + ("-entrees" if spec["kind"] == "ent" else "-stocks")
        ind = {
            "id": spec["id"],
            "title": spec["title"],
            "description": spec["description"],
            "unit": spec["unit"],
            "topic": spec["topic"],
            "national": True,
            "source_label": "DARES, POEM series " + ", ".join(c for c, _, _ in raws),
            "comparability": spec["note"] + COMMON_NOTE,
            "dims": [{"key": "scope", "label": "Scope",
                      "values": {c: l for c, l, _, _ in BREAKDOWNS}, "default": "TOTAL"}],
            "flags": {"p": "provisional (DARES revises recent months)"},
            "provenance": provenance(
                publisher=PUBLISHER, dataset_code=total_code,
                source_url=PAGE.format(slug=page_slug),
                api_url=EXPORT.format(code=total_code) + " (and the private/public/secondary/higher series: "
                        + ", ".join(c for c, _, _ in raws[1:]) + ")",
                licence=LICENCE,
                citation="DARES, Système d'information sur l'apprentissage — POEM series "
                         + ", ".join(c for c, _, _ in raws) + ".",
                source_updated=max(updated) if updated else None,
                raw=total_raw, raw_sha256=total_digest),
            "series": series,
        }
        written.append(rel(write_indicator(ind)))
    return written
