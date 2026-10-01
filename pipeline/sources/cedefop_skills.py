"""Cedefop skills-intelligence datasets (the /files/ downloads in Cedefop's Datasets section).

Like cedefop_files.py: Cedefop's HTML pages refuse automated requests, but the
/files/ downloads work. Each file is downloaded from its current URL (listed in
FILES; names change with each release) and stored under a stable raw name in
data/raw/cedefop/. If a download fails, the most recent stored copy is used.

Produces:
  indicators/cedefop-esi                    European Skills Index 2024 dataset (releases 2017–2024)
  indicators/cedefop-esi-editions           ESI 2018, 2020 and 2022 editions as published
  indicators/cedefop-clssi                  Labour and skills shortage index by occupation (ISCO 2-digit)
  indicators/cedefop-stas-employment        STAS short-term employment projections by occupation (ISCO 1-digit)
  indicators/cedefop-stas-employment-growth STAS projected annual employment growth by occupation

Not turned into indicators (no country data): the OJA imbalance CSV (EU27 level,
ISCO 4-digit occupations only) and the two IVET attractiveness frameworks
(conceptual inventories of factors and candidate indicators, no values).
"""

from __future__ import annotations

import io
import re
import warnings

from ..common import (FetchError, fetch, geo_code, latest_raw, num, provenance, rel,
                      save_raw, sha256, slug, write_indicator, warn)

SOURCE = {
    "id": "cedefop-skills",
    "name": "Cedefop — Skills intelligence datasets (ESI, CLSSI, STAS)",
    "publisher": "Cedefop",
    "homepage": "https://www.cedefop.europa.eu/en/datasets",
    "description": "Cedefop's downloadable skills-intelligence datasets: the European Skills Index (2018–2024 editions), the Cedefop labour and skills shortage index (CLSSI) by occupation, and the short-term anticipation of skills trends and VET demand (STAS) employment projections.",
    "access": "file",
    "browser_cors": False,
    "licence": "CC-BY-4.0",
    "cadence": "ESI biennial; STAS twice a year (after AMECO spring/autumn releases); CLSSI irregular",
    "secret": None,
    "outputs": ["indicators/cedefop-esi", "indicators/cedefop-esi-editions", "indicators/cedefop-clssi",
                "indicators/cedefop-stas-employment", "indicators/cedefop-stas-employment-growth"],
}

# Some Cedefop workbooks use an Excel extension openpyxl does not know; harmless.
warnings.filterwarnings("ignore", message="Unknown extension is not supported", module="openpyxl")

BASE = "https://www.cedefop.europa.eu/files/"
ESI_LANDING = "https://www.cedefop.europa.eu/en/datasets/dataset-european-skills-index-esi"

FILES = {
    "esi-2024": {"file": "esi_dataset_2024_formatted.xlsx", "raw_name": "esi-2024.xlsx", "landing": ESI_LANDING},
    "esi-2022": {"file": "esi_2022_scores.xlsx", "raw_name": "esi-2022.xlsx", "landing": ESI_LANDING},
    "esi-2020": {"file": "esi_2020_scores.xlsx", "raw_name": "esi-2020.xlsx", "landing": ESI_LANDING},
    "esi-2018": {"file": "esi_2018_scores.xlsx", "raw_name": "esi-2018.xlsx", "landing": ESI_LANDING},
    "clssi": {
        "file": "2026_cedefop_labour_skills_shortage_index_clssi_dataset.xlsx",
        "raw_name": "clssi.xlsx",
        "landing": "https://www.cedefop.europa.eu/en/datasets/labour-skills-shortage-index",
    },
    "stas": {
        "file": "stas_dataset_release_2026-08.xlsx",
        "raw_name": "stas.xlsx",
        "landing": "https://www.cedefop.europa.eu/en/datasets/stas",
        "doi": "10.2906/467749508762302",
    },
}


def download(key: str) -> tuple[bytes, str, object, object]:
    """Fetch a Cedefop file; fall back to the latest stored original. Returns (body, url, raw path, sha256)."""
    f = FILES[key]
    url = BASE + f["file"]
    try:
        body = fetch(url, timeout=180)
        if body[:2] == b"PK":  # xlsx is a zip
            raw, digest = save_raw("cedefop", f["raw_name"], body)
            return body, url, raw, digest
    except FetchError:
        pass
    stem, _, ext = f["raw_name"].rpartition(".")
    path = latest_raw("cedefop", f"*-{stem}.{ext}")
    if not path:
        raise FetchError(f"Could not download {url} and no stored copy in data/raw/cedefop/")
    warn(f"download refused or failed ({url}); used the stored copy {rel(path)}, which may be out of date")
    body = path.read_bytes()
    return body, url, path, sha256(body)


def _wb(body: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)


def _s(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def _sentences(parts) -> str:
    """Join note lines from a sheet into sentences."""
    return " ".join(p if p.endswith((".", ":")) else p + "." for p in parts)


ESI_LICENCE_NOTE = ("Licence CC BY 4.0 as stated on the Cedefop dataset landing page "
                    "(the workbook itself carries no licence statement).")

# ------------------------------------------------------------------ ESI ----

ESI_PILLARS = {
    "Skills Development": ["Basic education", "Training and other education"],
    "Skills Activation": ["Transition to work", "Labour market participation"],
    "Skills Matching": ["Skills utilisation", "Skills mismatch"],
}
ESI_SUBPILLARS = {sp: p for p, sps in ESI_PILLARS.items() for sp in sps}


def _esi_component(name: str) -> tuple[str, str]:
    """Map an ESI column name to (dims key, label)."""
    name = _s(name)
    if name in ("European Skills Index", "Index"):
        return "index", "European Skills Index (overall)"
    if name in ESI_PILLARS:
        return slug(name), f"Pillar: {name}"
    if name in ESI_SUBPILLARS:
        return slug(name), f"Sub-pillar: {name} ({ESI_SUBPILLARS[name]})"
    return slug(name), f"Indicator score: {name}"


def run_esi_2024() -> str:
    body, url, raw, digest = download("esi-2024")
    wb = _wb(body)
    notes = _sentences(_s(c) for r in wb["Notes"].iter_rows(values_only=True) for c in r
                       if _s(c) and "@" not in _s(c) and _s(c) != "NOTES" and not _s(c).startswith("Data set for"))
    rows = list(wb["Scores"].iter_rows(values_only=True))
    hdr = [_s(h) for h in rows[0]]
    ix = {h: i for i, h in enumerate(hdr) if h}
    comp_cols = [(i, h) for i, h in enumerate(hdr) if i > ix["Index Rank"] and h]
    order = (["index"] + [slug(p) for p in ESI_PILLARS] + [slug(s) for s in ESI_SUBPILLARS])
    values = {}
    series = []
    for r in rows[1:]:
        geo = geo_code(_s(r[ix["Country Code"]]))
        year = _s(r[ix["Release Year"]])
        if not geo or not re.fullmatch(r"20\d\d", year):
            continue
        for i, h in comp_cols:
            key, label = _esi_component(h)
            values[key] = label
            v = num(r[i])
            if v is not None:
                series.append({"geo": geo, "time": year, "value": v, "dims": {"component": key}})
    keys = sorted(values, key=lambda k: (order.index(k) if k in order else len(order), values[k]))
    ind = {
        "id": "cedefop-esi",
        "title": "European Skills Index (ESI), 2024 methodology",
        "description": (
            "Cedefop's European Skills Index measures how well each country's skills system performs, "
            "as a composite of 15 indicators grouped into three pillars — skills development (basic "
            "education; training and other education), skills activation (transition to work; labour "
            "market participation) and skills matching (skills utilisation; skills mismatch). Scores are "
            "normalised from 0 to 100, where 100 is the 'frontier' of best achievable performance. "
            "The ESI 2024 dataset recomputes every release from 2017 to 2024 with the 2024 methodology. "
            "Time is the ESI release year; the release year equals the data year plus 2 "
            "(e.g. ESI 2024 uses principally 2022 data)."),
        "unit": "score (0–100)",
        "topic": "Skills intelligence",
        "source_label": "Cedefop, European Skills Index 2024 dataset (Scores sheet)",
        "comparability": (
            "All releases 2017–2024 in this series use the same (2024) methodology, so they are comparable "
            "over time and across countries. They are not comparable with the scores originally published in "
            "the ESI 2018, 2020 and 2022 editions (see cedefop-esi-editions). Indicator scores are normalised "
            "distances to the frontier, not raw values; the underlying raw data (before imputation) are in "
            "the source workbook's RawData sheet. No EU27 aggregate is published. Cedefop notes: " + notes),
        "dims": [{"key": "component", "label": "Index, pillar, sub-pillar or indicator",
                  "values": {k: values[k] for k in keys}, "default": "index"}],
        "provenance": provenance(
            publisher="Cedefop", dataset_code="ESI 2024", source_url=FILES["esi-2024"]["landing"],
            api_url=url, licence="CC-BY-4.0",
            citation="Cedefop (2024). European Skills Index 2024 [Data set]. "
                     "https://www.cedefop.europa.eu/en/datasets/dataset-european-skills-index-esi. " + ESI_LICENCE_NOTE,
            source_updated="ESI 2024 (dataset page dated 23 February 2024)", raw=raw, raw_sha256=digest),
        "series": series,
    }
    return rel(write_indicator(ind))


ESI_EDITION_KEYS = ["esi-2018", "esi-2020", "esi-2022"]


def parse_esi_edition(body: bytes) -> tuple[list[dict], str]:
    """Composite scores (index, pillars, sub-pillars) from an ESI 2018/2020/2022 'Scores & Ranks' sheet.

    Layout: row 0 = block label (merged), row 1 = Score/Rank (merged), row 2 = data year,
    then one row per country code. Returns observations with 'component' set, and the notes text.
    """
    wb = _wb(body)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    width = max(len(r) for r in rows[:3])
    block = kind = None
    cols = []
    for j in range(1, width):
        lab = _s(rows[0][j]) if j < len(rows[0]) else ""
        if lab:
            block = lab
        k = _s(rows[1][j]) if j < len(rows[1]) else ""
        if k:
            kind = k
        year = _s(rows[2][j]) if j < len(rows[2]) else ""
        if not (block and kind == "Score" and re.fullmatch(r"20\d\d", year)):
            continue
        name = re.sub(r"^(Pillar|Sub-pillar):\s*", "", block)
        if name == "Index" or name in ESI_PILLARS or name in ESI_SUBPILLARS:
            cols.append((j, _esi_component(name)[0], year))
    obs = []
    notes = []
    for r in rows[3:]:
        first = _s(r[0]) if r else ""
        if first == "Notes:" or notes:
            notes += [_s(c) for c in r if _s(c) and _s(c) != "Notes:"]
            continue
        geo = geo_code(first)
        if not geo:
            continue
        for j, key, year in cols:
            v = num(r[j]) if j < len(r) else None
            if v is not None:
                obs.append({"geo": geo, "time": year, "value": v, "component": key})
    # Some editions store scores on 0–1, others (the 2018 overall index) on 0–100: put all on 0–100.
    for key in {o["component"] for o in obs}:
        part = [o for o in obs if o["component"] == key]
        if max(o["value"] for o in part) <= 1.0:
            for o in part:
                o["value"] *= 100
    notes = [n for n in notes if "@" not in n]  # drop the contact line
    return obs, _sentences(notes)


def run_esi_editions() -> str:
    series = []
    raws = []
    urls = []
    notes = ""
    for key in ESI_EDITION_KEYS:
        body, url, raw, digest = download(key)
        edition = key.rsplit("-", 1)[1]
        obs, notes = parse_esi_edition(body)
        for o in obs:
            series.append({"geo": o["geo"], "time": o["time"], "value": o["value"],
                           "dims": {"edition": edition, "component": o["component"]}})
        if key != ESI_EDITION_KEYS[-1]:
            raws.append(f"{rel(raw)} (sha256 {digest})")
        urls.append(url)
    comps = ["index"] + [slug(p) for p in ESI_PILLARS] + [slug(s) for s in ESI_SUBPILLARS]
    present = {o["dims"]["component"] for o in series}
    last = ESI_EDITION_KEYS[-1]
    ind = {
        "id": "cedefop-esi-editions",
        "title": "European Skills Index (ESI) — 2018, 2020 and 2022 editions as published",
        "description": (
            "Overall index, pillar and sub-pillar scores of Cedefop's European Skills Index as published in "
            "the 2018, 2020 and 2022 editions. Each edition covers three data years (2018 edition: 2014–2016; "
            "2020 edition: 2016–2018; 2022 edition: 2018–2020). Time is the data year, i.e. the year the "
            "underlying indicators refer to, not the release year. Scores are distances to the 'frontier' of "
            "best achievable performance, shown on a 0–100 scale."),
        "unit": "score (0–100)",
        "topic": "Skills intelligence",
        "source_label": "Cedefop, European Skills Index edition score files (Scores & Ranks sheets)",
        "comparability": (
            "Do not compare across editions: each edition re-estimated all years with its own methodology, "
            "indicator set and frontier, so the same country and data year has different scores in different "
            "editions. Within an edition, years and countries are comparable. Country coverage differs "
            "(2018: EU28; 2020 and 2022: EU27 + UK, CH, IS, NO). The 2018 overall index was published on a "
            "0–100 scale and all other scores on a 0–1 scale; apprentix.eu multiplied the 0–1 scores by 100. "
            "The 2022 edition file gives scores rounded to two decimals. For a consistent time series use "
            "cedefop-esi (ESI 2024 dataset, releases 2017–2024). Cedefop notes: " + notes),
        "dims": [
            {"key": "edition", "label": "ESI edition",
             "values": {e.rsplit("-", 1)[1]: f"ESI {e.rsplit('-', 1)[1]}" for e in ESI_EDITION_KEYS},
             "default": last.rsplit("-", 1)[1]},
            {"key": "component", "label": "Index, pillar or sub-pillar",
             "values": {k: _label_for(k) for k in comps if k in present}, "default": "index"},
        ],
        "provenance": provenance(
            publisher="Cedefop", dataset_code="ESI 2018/2020/2022", source_url=ESI_LANDING,
            api_url=" ; ".join(urls), licence="CC-BY-4.0",
            citation="Cedefop. European Skills Index, 2018, 2020 and 2022 editions [Data sets]. "
                     "https://www.cedefop.europa.eu/en/datasets/dataset-european-skills-index-esi. "
                     + ESI_LICENCE_NOTE + " Other raw files: " + "; ".join(raws),
            source_updated="ESI 2018, 2020 and 2022 editions", raw=raw, raw_sha256=digest),
        "series": series,
    }
    return rel(write_indicator(ind))


def _label_for(key: str) -> str:
    for name in ["Index", *ESI_PILLARS, *ESI_SUBPILLARS]:
        k, label = _esi_component(name)
        if k == key:
            return label
    return key


# ---------------------------------------------------------------- CLSSI ----

ISCO08_2D = {
    "Chief executives, senior officials and legislators": "11",
    "Administrative and commercial managers": "12",
    "Production and specialised services managers": "13",
    "Hospitality, retail and other services managers": "14",
    "Science and engineering professionals": "21",
    "Health professionals": "22",
    "Teaching professionals": "23",
    "Business and administration professionals": "24",
    "Information and communications technology professionals": "25",
    "Legal, social and cultural professionals": "26",
    "Science and engineering associate professionals": "31",
    "Health associate professionals": "32",
    "Business and administration associate professionals": "33",
    "Legal, social, cultural and related associate professionals": "34",
    "Information and communications technicians": "35",
    "General and keyboard clerks": "41",
    "Customer services clerks": "42",
    "Numerical and material recording clerks": "43",
    "Other clerical support workers": "44",
    "Personal service workers": "51",
    "Sales workers": "52",
    "Personal care workers": "53",
    "Protective services workers": "54",
    "Market-oriented skilled agricultural workers": "61",
    "Market-oriented skilled forestry, fishery and hunting workers": "62",
    "Subsistence farmers, fishers, hunters and gatherers": "63",
    "Building and related trades workers, excluding electricians": "71",
    "Metal, machinery and related trades workers": "72",
    "Handicraft and printing workers": "73",
    "Electrical and electronic trades workers": "74",
    "Food processing, wood working, garment and other craft and related trades": "75",
    "Food processing, wood working, garment and other craft and related trades workers": "75",
    "Stationary plant and machine operators": "81",
    "Assemblers": "82",
    "Drivers and mobile plant operators": "83",
    "Cleaners and helpers": "91",
    "Agricultural, forestry and fishery labourers": "92",
    "Labourers in mining, construction, manufacturing and transport": "93",
    "Food preparation assistants": "94",
    "Street and related sales and service workers": "95",
    "Refuse workers and other elementary workers": "96",
}


def run_clssi() -> str:
    body, url, raw, digest = download("clssi")
    wb = _wb(body)
    series = []
    occ = {}
    groups = {}
    for ws in wb.worksheets:
        geo = geo_code(ws.title.strip())
        if not geo:
            continue
        rows = list(ws.iter_rows(values_only=True))
        hdr = [_s(h).lower() for h in rows[0]]
        i_name = next(i for i, h in enumerate(hdr) if h.startswith("occupation group"))
        i_grp = next(i for i, h in enumerate(hdr) if h.startswith("main occupation"))
        i_idx = next(i for i, h in enumerate(hdr) if h.startswith("labour shortage ind"))
        i_comp = [i for i, h in enumerate(hdr) if re.fullmatch(r"lsi\d", h)]
        for r in rows[1:]:
            name = _s(r[i_name]) if i_name < len(r) else ""
            if not name:
                continue
            code = ISCO08_2D.get(name)
            if code is None:
                raise ValueError(f"CLSSI: unknown occupation {name!r} in sheet {ws.title}")
            occ[code] = f"{code} {name}"
            groups[code] = _s(r[i_grp])
            v = num(r[i_idx])
            if v is None:
                continue
            comps = [num(r[i]) for i in i_comp]
            flag = "m" if any(c is None for c in comps) else None
            series.append({"geo": geo, "time": "2026", "value": v, "dims": {"occupation": code},
                           **({"flag": flag} if flag else {})})
    group_note = "; ".join(
        f"{g}: ISCO {', '.join(sorted(c for c in groups if groups[c] == g))}"
        for g in dict.fromkeys(groups[c] for c in sorted(groups)))
    ind = {
        "id": "cedefop-clssi",
        "title": "Cedefop labour and skills shortage index (CLSSI) by occupation",
        "description": (
            "Cedefop's labour and skills shortage index for 38 occupation groups (ISCO-08 2-digit), by country. "
            "The index is the average of three component shortage indicators (LSI1, LSI2, LSI3 in the source), "
            "each scored from 1 to 4; higher values signal stronger labour and skills shortages in that "
            "occupation. Time is the release year of the dataset file (2026)."),
        "unit": "index (1–4)",
        "topic": "Skills intelligence",
        "source_label": "Cedefop labour and skills shortage index (CLSSI), 2026 dataset",
        "comparability": (
            "Scores are on an ordinal 1–4 scale; see the Cedefop dataset page for how the three component "
            "indicators are defined and banded. The workbook does not state "
            "the reference period of the underlying data (the dataset page dates from November 2024, the "
            "file from 2026). Where one component indicator is missing the index is the mean of the two "
            "available (flag 'm'). ISCO-08 codes were added by apprentix.eu from the occupation titles; "
            "occupation groups 63 and 95 are not covered. Cedefop's broad groups: " + group_note + "."),
        "dims": [{"key": "occupation", "label": "Occupation (ISCO-08 2-digit)",
                  "values": dict(sorted(occ.items())), "default": "71"}],
        "flags": {"m": "one of the three component indicators is missing; index averages the other two"},
        "provenance": provenance(
            publisher="Cedefop", dataset_code="CLSSI 2026", source_url=FILES["clssi"]["landing"],
            api_url=url, licence="CC-BY-4.0",
            citation="Cedefop (2026). Cedefop labour and skills shortage index (CLSSI) [Data set]. "
                     "https://www.cedefop.europa.eu/en/datasets/labour-skills-shortage-index. Licence CC BY 4.0 "
                     "as stated on the Cedefop dataset landing page (the workbook carries no licence statement).",
            source_updated="2026 dataset", raw=raw, raw_sha256=digest),
        "series": series,
    }
    return rel(write_indicator(ind))


# ----------------------------------------------------------------- STAS ----

def parse_stas_sheet(ws) -> tuple[dict, list[tuple]]:
    """Return (meta, rows) for a STAS sheet: 5 'key/value' lines, a header row, then data."""
    rows = list(ws.iter_rows(values_only=True))
    meta = {_s(r[0]): _s(r[1]) for r in rows[:5]}
    hdr = rows[5]
    years = [(j, str(int(h)) if isinstance(h, (int, float)) else _s(h)) for j, h in enumerate(hdr)
             if re.fullmatch(r"20\d\d", str(int(h)) if isinstance(h, (int, float)) else _s(h))]
    out = []
    for r in rows[6:]:
        geo = geo_code(_s(r[2]))
        if not geo:
            continue
        code = _s(r[4])
        code = "TOTAL" if code == "TOT" else str(int(float(code)))
        for j, y in years:
            v = num(r[j]) if j < len(r) else None
            if v is not None:
                out.append((geo, y, code, _s(r[3]), _s(r[0]), v))
    return meta, out


def run_stas() -> list[str]:
    body, url, raw, digest = download("stas")
    wb = _wb(body)
    readme = {_s(r[0]).rstrip(":"): _s(r[1]) for r in wb["Readme"].iter_rows(values_only=True) if r and _s(r[0])}
    f = FILES["stas"]
    citation = (f"Cedefop (2026). Short-term anticipation of skills trends and VET demand (STAS), "
                f"{readme.get('Version', '')} release [Data set]. https://doi.org/{f['doi']}. "
                f"© {readme.get('Copyright', 'Cedefop')}; licence CC BY 4.0 per the Cedefop dataset page and "
                f"the data.europa.eu record (the workbook states copyright only). {readme.get('ARES number', '')}")
    common_comp = (
        "Short-term projections, not observations: STAS projects employment by occupation from EU-LFS "
        "quarterly and annual data and European job vacancy statistics, aligned to DG ECFIN's AMECO forecast of "
        "total employment growth (scenario 'Aligned_forecast_Ameco'). Updated twice a year, so values for the "
        "same year change between releases. Occupation titles in the 1-digit tables follow the source "
        "(ISCO major groups 1–9). EU27 is published by Cedefop. EU Member States only.")
    written = []
    for sheet, ind_id, title, unit, scale, desc in [
        ("ameco_1d", "cedefop-stas-employment",
         "Projected employment by occupation (STAS short-term projections)", "thousand persons", 1,
         "Cedefop's short-term projection of the number of people employed, by ISCO major occupation group "
         "(1-digit) and in total, for the current and next year."),
        ("ameco_1d_%", "cedefop-stas-employment-growth",
         "Projected annual employment growth by occupation (STAS short-term projections)", "%", 100,
         "Cedefop's short-term projection of the annual growth rate of employment (percentage change on the "
         "previous year), by ISCO major occupation group (1-digit) and in total."),
    ]:
        meta, rows = parse_stas_sheet(wb[sheet])
        occ = {}
        series = []
        for geo, y, code, name, scen, v in rows:
            occ[code] = "Total" if code == "TOTAL" else f"{code} {name}"
            series.append({"geo": geo, "time": y, "value": v * scale, "flag": "f",
                           "dims": {"occupation": code}})
        values = {"TOTAL": occ.pop("TOTAL")} if "TOTAL" in occ else {}
        values.update(dict(sorted(occ.items(), key=lambda kv: int(kv[0]))))
        ind = {
            "id": ind_id,
            "title": title,
            "description": desc + f" Release: {readme.get('Description', '')}.",
            "unit": unit,
            "topic": "Skills intelligence",
            "source_label": f"Cedefop STAS, sheet {sheet} ({meta.get('Desc', '')}, {meta.get('Presented as', '')})",
            "comparability": common_comp + (" Growth rates are published as fractions and were multiplied by "
                                            "100 by apprentix.eu." if scale == 100 else ""),
            "dims": [{"key": "occupation", "label": "Occupation (ISCO major group)",
                      "values": values, "default": "TOTAL"}],
            "flags": {"f": "forecast (short-term projection)"},
            "provenance": provenance(
                publisher="Cedefop", dataset_code=f"STAS {readme.get('Version', '')} ({sheet})",
                source_url=f["landing"], api_url=url, licence="CC-BY-4.0", citation=citation,
                source_updated=readme.get("Version"), raw=raw, raw_sha256=digest),
            "series": series,
        }
        written.append(rel(write_indicator(ind)))
    return written


def run() -> list[str]:
    return [run_esi_2024(), run_esi_editions(), run_clssi(), *run_stas()]
