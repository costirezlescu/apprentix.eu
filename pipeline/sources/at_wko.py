"""Austria — WKO Lehrlingsstatistik (Austrian Economic Chambers apprenticeship statistics).

Landing page: https://www.wko.at/zahlen-daten-fakten/daten-lehrlingsstatistik
The Excel files there have stable, year-less names under /oe/statistik/jahrbuch/ and
are replaced in place each year (reference date 31 December). Each workbook carries
an 'olap_*' sheet with the full annual pivot (2002 onwards), which is what we parse;
the presentation sheet of ll-zeitr-sp-bdl.xlsx adds the benchmark years 1980/1990/2000.

!! Reuse terms: wko.at's legal notice (Offenlegung, https://www.wko.at/offenlegung-oesterreich)
states the content is protected by copyright, for personal use only, and that storage in
databases, reproduction, commercial use and passing on to third parties — also in parts
or in revised form — is prohibited without the consent of the organisation concerned.
No open licence was found for the Lehrlingsstatistik. Obtain WKO's permission before
publishing these outputs.

Produces (geo = AT, national figures):
  indicators/at-wko-apprentices        apprentices at 31.12 by Sparte (sector), apprenticeship year, sex
  indicators/at-wko-apprentices-land   apprentices at 31.12 by Bundesland, apprenticeship year, sex
"""

from __future__ import annotations

import io
import re

from ..common import FetchError, fetch, latest_raw, provenance, rel, save_raw, write_indicator

BASE = "https://www.wko.at/oe/statistik/jahrbuch/"
LANDING = "https://www.wko.at/zahlen-daten-fakten/daten-lehrlingsstatistik"
NOTES = "https://www.wko.at/oe/statistik/lehrling/erlaeuterungen.pdf"

LICENCE = ("unverified — no open licence. wko.at legal notice: 'Der Inhalt dieser Homepage ist "
           "urheberrechtlich geschützt. Die Informationen sind nur für die persönliche Verwendung "
           "bestimmt. Jede weitergehende Nutzung insbesondere die Speicherung in Datenbanken, "
           "Vervielfältigung und jede Form von gewerblicher Nutzung sowie die Weitergabe an Dritte - "
           "auch in Teilen oder in überarbeiteter Form - ohne Zustimmung der jeweiligen Organisation "
           "ist untersagt.' Permission from WKO is needed before republishing.")
LICENCE_URL = "https://www.wko.at/offenlegung-oesterreich"

SOURCE = {
    "id": "at-wko",
    "name": "Austria — WKO Lehrlingsstatistik",
    "publisher": "Wirtschaftskammer Österreich (WKO), Abteilung für Statistik",
    "homepage": LANDING,
    "description": "Apprentices (Lehrlinge) in Austria at 31 December by sector (Sparte), federal state, apprenticeship year and sex, from the apprenticeship offices of the Economic Chambers. National series, 2002 onwards plus 1980/1990/2000 benchmarks.",
    "access": "file",
    "browser_cors": None,
    "licence": "unverified — wko.at terms reserve all rights (personal use only; storage in databases and redistribution need WKO consent)",
    "cadence": "annual (published in January–February for 31 December of the previous year)",
    "secret": None,
    # Not published until the publisher agrees; see pipeline/run.py.
    "permission_needed": "WKO's terms forbid storing the data in databases or passing it on without WKO's consent. Ask WKO (Abteilung für Statistik) for permission before enabling.",
    "reuse_note": "WKO consent required before republishing (see wko.at Offenlegung).",
    "outputs": ["indicators/at-wko-apprentices", "indicators/at-wko-apprentices-land"],
}

FILES = {
    "sector": ("ll-sp-lj-gesch.xlsx", "olap_sp_lj_gesch"),
    "land": ("ll-bdl-lj-gesch.xlsx", "olap_bld_lj_gesch"),
    "longrun": ("ll-zeitr-sp-bdl.xlsx", "LL_Zeitr_Sp_BDL"),
}

SECTORS = {
    "Gewerbe & Handwerk": ("GEWERBE", "Trade and crafts (Gewerbe und Handwerk)"),
    "Industrie": ("INDUSTRIE", "Industry"),
    "Handel": ("HANDEL", "Commerce (Handel)"),
    "Bank & Versicherung": ("BANK", "Banking and insurance"),
    "Transport & Verkehr": ("TRANSPORT", "Transport and traffic"),
    "Tourismus & Freizeitwirtschaft": ("TOURISMUS", "Tourism and leisure"),
    "Information & Consulting": ("INFO", "Information and consulting"),
    "Sonstige Lehrberechtigte": ("SONSTIGE", "Other training companies outside the Chamber (e.g. law firms, municipalities)"),
    "Überbetriebliche Lehrausbildung": ("UEBA", "Supra-company training (Überbetriebliche Lehrausbildung)"),
    "Gesamtergebnis": ("TOTAL", "Total"),
}
LAENDER = {
    "Burgenland": ("AT11", "Burgenland"), "Kärnten": ("AT21", "Carinthia (Kärnten)"),
    "Niederösterreich": ("AT12", "Lower Austria (Niederösterreich)"),
    "Oberösterreich": ("AT31", "Upper Austria (Oberösterreich)"), "Salzburg": ("AT32", "Salzburg"),
    "Steiermark": ("AT22", "Styria (Steiermark)"), "Tirol": ("AT33", "Tyrol (Tirol)"),
    "Vorarlberg": ("AT34", "Vorarlberg"), "Wien": ("AT13", "Vienna (Wien)"),
    "Gesamtergebnis": ("AT", "Austria"),
}
SEX = {"männlich": "M", "weiblich": "F", "divers": "X"}
SEX_LABELS = {"T": "Total", "M": "Males", "F": "Females", "X": "Diverse / other (recorded since 2023)"}
LJ_LABELS = {"TOTAL": "All apprenticeship years", "Y1": "1st year", "Y2": "2nd year",
             "Y3": "3rd year", "Y4": "4th year"}

HEADER = re.compile(r"^(\d{4})_(\d)\.LJ_(männlich|weiblich|divers)$|^(\d{4})_ins_(männlich|weiblich|divers)$")

COMPARABILITY = (
    "National series from the WKO apprenticeship statistics (administrative records of the Chambers' "
    "apprenticeship offices). Stock of apprentices with a registered apprenticeship contract (Lehrvertrag) "
    "on 31 December, all sectors including supra-company training (ÜBA) and training under § 8b BAG "
    "(extended apprenticeships and partial qualifications). Dual apprenticeship only — school-based VET "
    "(BMS/BHS) is not included. 'Total' sex = male + female + diverse. Do not add to other countries' figures."
)


def download(name: str) -> tuple[bytes, object, str]:
    url = BASE + name
    try:
        body = fetch(url, timeout=120, polite_delay=1.0)
        if body[:2] != b"PK":
            raise FetchError(f"not an xlsx: {url}")
        raw, digest = save_raw("at-wko", name, body)
    except FetchError:
        raw = latest_raw("at-wko", "*-" + name)
        if raw is None:
            raise
        body = raw.read_bytes()
        from ..common import sha256
        digest = sha256(body)
    return body, raw, digest


def parse_olap(body: bytes, sheet: str, rows_map: dict) -> dict:
    """Return {(row_code, year, lj, sex): value} from a WKO pivot ('olap') sheet.

    Header rows have 1 in column A and 'YYYY_n.LJ_sex' / 'YYYY_ins_sex' labels;
    the data rows follow until the next blank row. Empty pivot cells are zeros.
    """
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)[sheet]
    out, cols = {}, None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == 1:
            cols = {}
            for i, h in enumerate(row):
                m = HEADER.match(str(h).strip()) if h is not None else None
                if m:
                    if m.group(1):
                        cols[i] = (m.group(1), "Y" + m.group(2), SEX[m.group(3)])
                    else:
                        cols[i] = (m.group(4), "TOTAL", SEX[m.group(5)])
            continue
        if not row or all(v is None for v in row[:3]):
            cols = None
            continue
        label = str(row[1]).strip() if len(row) > 1 and row[1] is not None else None
        if cols and label in rows_map:
            code = rows_map[label][0]
            for i, key in cols.items():
                v = row[i] if i < len(row) else None
                out[(code, *key)] = float(v) if isinstance(v, (int, float)) else 0.0
    if not out:
        raise ValueError(f"no data parsed from sheet {sheet}")
    return out


def longrun(body: bytes) -> dict:
    """Benchmark years before 2002 (1980, 1990, 2000) from the presentation sheet: {(code, year): value}."""
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)[FILES["longrun"][1]]
    out, years = {}, None
    names = {"INSGESAMT": "TOTAL", "ÖSTERREICH": "AT", **{k: v[0] for k, v in LAENDER.items() if k != "Gesamtergebnis"}}
    for row in ws.iter_rows(values_only=True):
        first = str(row[0]).strip() if row and row[0] is not None else ""
        if first in ("Sparte", "Bundesland"):
            years = {i: int(v) for i, v in enumerate(row) if isinstance(v, (int, float)) and 1900 < v < 2002}
            continue
        if years and first in names:
            for i, y in years.items():
                if isinstance(row[i], (int, float)):
                    out[(names[first], str(y))] = float(row[i])
    return out


def series_from(cube: dict, dim_key: str) -> list[dict]:
    """Add sex totals (M+F+X) and return observations."""
    totals = {}
    for (code, year, lj, sex), v in cube.items():
        totals[(code, year, lj)] = totals.get((code, year, lj), 0.0) + v
    obs = [{"geo": "AT", "time": y, "value": v, "dims": {dim_key: c, "year_of_training": lj, "sex": s}}
           for (c, y, lj, s), v in cube.items()
           if not (s == "X" and v == 0 and int(y) < 2023)]
    obs += [{"geo": "AT", "time": y, "value": v, "dims": {dim_key: c, "year_of_training": lj, "sex": "T"}}
            for (c, y, lj), v in totals.items()]
    return obs


def check(cube: dict, label: str) -> None:
    """Sanity check: apprenticeship years add up to the all-years block."""
    sums, tot = {}, {}
    for (c, y, lj, s), v in cube.items():
        if lj == "TOTAL":
            tot[(c, y, s)] = v
        else:
            sums[(c, y, s)] = sums.get((c, y, s), 0) + v
    bad = [k for k in tot if abs(tot[k] - sums.get(k, 0)) > 2]
    if len(bad) > len(tot) * 0.05:
        raise ValueError(f"{label}: apprenticeship-year breakdown does not add up for {bad[:5]}")


def indicator(id_, title, description, dim_key, dim_label, dim_values, obs, file, raw, digest, extra_note=""):
    years = sorted({o["time"] for o in obs})
    return {
        "id": id_,
        "title": title,
        "description": description,
        "unit": "apprentices",
        "topic": "Participation",
        "national": True,
        "source_label": f"WKO Lehrlingsstatistik, Stichtag 31.12. ({years[0]}–{years[-1]})",
        "comparability": COMPARABILITY + extra_note,
        "dims": [
            {"key": dim_key, "label": dim_label, "values": dim_values, "default": "TOTAL" if "TOTAL" in dim_values else "AT"},
            {"key": "year_of_training", "label": "Apprenticeship year", "values": LJ_LABELS, "default": "TOTAL"},
            {"key": "sex", "label": "Sex", "values": SEX_LABELS, "default": "T"},
        ],
        "provenance": {
            **provenance(publisher="Wirtschaftskammer Österreich (WKO)", dataset_code=file,
                         source_url=LANDING, api_url=BASE + file, licence=LICENCE,
                         citation=f"WKO, Lehrlingsstatistik (Stichtag 31.12.), {file}. Quelle: LEHRLINGSSTATISTIK, Wirtschaftskammern Österreichs.",
                         raw=raw, raw_sha256=digest),
            "licence_url": LICENCE_URL,
        },
        "series": obs,
    }


def run() -> list[str]:
    written = []
    lr_body, _, _ = download(FILES["longrun"][0])
    bench = longrun(lr_body)
    bench_note = (" Values for 1980, 1990 and 2000 are national/Land totals only, taken from WKO's long-run table"
                  " (ll-zeitr-sp-bdl.xlsx); annual breakdowns start in 2002.")

    # By sector (Sparte).
    file, sheet = FILES["sector"]
    body, raw, digest = download(file)
    cube = parse_olap(body, sheet, SECTORS)
    check(cube, file)
    obs = series_from(cube, "sector")
    obs += [{"geo": "AT", "time": y, "value": v, "dims": {"sector": "TOTAL", "year_of_training": "TOTAL", "sex": "T"}}
            for (c, y), v in bench.items() if c == "TOTAL"]
    written.append(rel(write_indicator(indicator(
        "at-wko-apprentices", "Apprentices by sector (Austria)",
        "Number of apprentices (Lehrlinge) in Austria on 31 December, by WKO sector (Sparte), apprenticeship year and sex.",
        "sector", "Sector (Sparte)", {v[0]: v[1] for v in SECTORS.values()}, obs, file, raw, digest,
        bench_note + " The sector is the Chamber section of the training company; ÜBA = training institutions commissioned by the public employment service (AMS)."))))

    # By federal state (Bundesland).
    file, sheet = FILES["land"]
    body, raw, digest = download(file)
    cube = parse_olap(body, sheet, LAENDER)
    check(cube, file)
    obs = series_from(cube, "land")
    obs += [{"geo": "AT", "time": y, "value": v, "dims": {"land": c, "year_of_training": "TOTAL", "sex": "T"}}
            for (c, y), v in bench.items() if c != "TOTAL"]
    written.append(rel(write_indicator(indicator(
        "at-wko-apprentices-land", "Apprentices by federal state (Austria)",
        "Number of apprentices (Lehrlinge) in Austria on 31 December, by federal state (Bundesland, NUTS 2 codes), apprenticeship year and sex. 'AT' is the national total.",
        "land", "Federal state (Bundesland)", {v[0]: v[1] for v in LAENDER.values()}, obs, file, raw, digest,
        bench_note))))
    return written
